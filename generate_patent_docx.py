import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_simple_patent_doc():
    doc = docx.Document()

    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Header Tag
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_badge = p_badge.add_run("INVENTION DISCLOSURE")
    r_badge.font.size = Pt(9)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("SchemaBridge AI — Patent Invention Disclosure")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Smart, High-Speed Data Translation Between Enterprise Systems")
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # Quick Info Box
    meta_data = [
        ("Project Name:", "SchemaBridge AI"),
        ("Inventor:", "Baljinder Singh"),
        ("Invention Title:", "System and Method for Fast, Secure Data Translation Between Enterprise Systems"),
        ("Core Idea:", "Uses AI only once to generate rules, then runs 100% pure fast code in production (zero AI in live data flow)")
    ]
    meta_table = doc.add_table(rows=len(meta_data), cols=2)
    for row_idx, (label, val) in enumerate(meta_data):
        c0 = meta_table.cell(row_idx, 0)
        c1 = meta_table.cell(row_idx, 1)
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)
        set_cell_background(c0, "F7FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80)
        set_cell_margins(c1, top=80, bottom=80)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)

    doc.add_paragraph()

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        return h

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            r_b = p.add_run("•  " + bold_prefix + ": ")
            r_b.font.bold = True
        else:
            r_b = p.add_run("•  ")
        p.add_run(text)
        return p

    # 1. The Core Idea
    add_heading_1("1. The Core Idea (In Plain English)")
    doc.add_paragraph(
        "Enterprise systems (like an old billing system and a modern cloud CRM) cannot talk to each other "
        "because their field names, date formats, and structures don't match."
    )
    doc.add_paragraph("SchemaBridge AI solves this with a smart two-phase design:")
    add_bullet("It uses AI only once at design-time to examine schemas and figure out translation rules.", "Phase 1 (Design Time)")
    add_bullet("It executes the translation in pure code (<5ms) without calling AI. Live customer data never touches an AI model.", "Phase 2 (Live Runtime)")

    # 2. Problems With Current Ways
    add_heading_1("2. Problems With Existing Solutions")
    add_bullet("Writing point-to-point integration code for hundreds of fields takes weeks and breaks every time an API updates.", "Manual Hand-Coding is Slow & Fragile")
    add_bullet("Calling an LLM for each live transaction takes 1 to 5 seconds. Production systems need answers in under 10 milliseconds.", "Sending Live Data to AI is Too Slow")
    add_bullet("LLMs can hallucinate, change output formats, or drop fields randomly, which causes data corruption in financial or healthcare records.", "AI is Unpredictable in Production")
    add_bullet("Sending customer data (names, emails, credit cards) to external AI models violates privacy laws (GDPR, HIPAA).", "Data Privacy & Compliance Risks")
    add_bullet("Paying token fees for millions of daily live API calls is far too expensive.", "High Token Costs")

    # 3. Simple Flow Diagram
    add_heading_1("3. Simple System Flow")
    
    diagram_lines = [
        " [System A Schema]                      [System B Schema]",
        "         │                                      │",
        "         ▼                                      ▼",
        "┌────────────────────────────────────────────────────────┐",
        "│             STEP 1: 5-LEVEL FIELD MATCHING             │",
        "│  Level 1: Exact Name Match (e.g., id -> id)           │",
        "│  Level 2: Formatted Match (e.g., user_name -> userName)│",
        "│  Level 3: Synonym Match (e.g., dob -> dateOfBirth)     │",
        "│  Level 4: Past History (Reuses previously saved rules) │",
        "│  Level 5: AI Matching (ONLY for leftover unknown fields│",
        "└───────────────────────────┬────────────────────────────┘",
        "                            │",
        "                            ▼",
        "┌────────────────────────────────────────────────────────┐",
        "│             STEP 2: HUMAN REVIEW & APPROVAL            │",
        "│  - High confidence rules: ready to go                  │",
        "│  - Low confidence rules: flagged for quick human check │",
        "└───────────────────────────┬────────────────────────────┘",
        "                            │ (Approved Translation Rules)",
        "                            ▼",
        "┌────────────────────────────────────────────────────────┐",
        "│             STEP 3: LIVE RUNTIME EXECUTION             │",
        "│  [Live System A Data] ──► [Pure Code Engine] ──► [System B]",
        "│                            - Renames fields            │",
        "│                            - Converts dates & types    │",
        "│                            - Flattens / nests objects  │",
        "│   * Speed: Under 5 milliseconds                        │",
        "│   * Privacy: Zero data sent to AI                      │",
        "│   * Accuracy: 100% deterministic (no hallucinations)   │",
        "└────────────────────────────────────────────────────────┘"
    ]
    p_diag = doc.add_paragraph()
    p_diag.paragraph_format.left_indent = Inches(0.15)
    p_diag.paragraph_format.space_before = Pt(4)
    p_diag.paragraph_format.space_after = Pt(8)
    r_diag = p_diag.add_run("\n".join(diagram_lines))
    r_diag.font.name = "Consolas"
    r_diag.font.size = Pt(8.5)
    r_diag.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    # 4. How the Steps Work
    add_heading_1("4. How the Steps Work (Step-by-Step)")
    add_bullet("Checks exact names, case formatting (user_id to userId), known business synonyms (dob to dateOfBirth), and past approved rules. Calls AI only for the few remaining fields that couldn't be resolved deterministically.", "Step 1: The 5-Level Matching Pipeline")
    add_bullet("Safe rules with high confidence scores pass automatically. Lower confidence matches are held in a review queue for a quick 1-click human check before going live.", "Step 2: Human-in-the-Loop Check")
    add_bullet("Live customer payloads run through pre-compiled algebraic rules (rename, date format, flatten, nest) in pure Java code (<5ms). Works offline even if the AI service goes down.", "Step 3: Zero-AI Runtime Engine")
    add_bullet("Cryptographically hashes schemas to catch changes immediately. Highlights exactly which rule is broken (BROKEN_PATH) before bad data hits production.", "Step 4: Automatic Schema Change Detection")

    # 5. Comparison Table
    add_heading_1("5. Comparison with Other Approaches")
    comp_table = doc.add_table(rows=7, cols=4)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["Feature", "Hand-Written Code", "Calling AI on Live Data", "SchemaBridge AI"]
    for c_idx, h_text in enumerate(c_headers):
        cell = comp_table.cell(0, c_idx)
        set_cell_background(cell, "2B6CB0")
        set_cell_margins(cell, top=100, bottom=100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)

    comp_rows = [
        ("Setup Time", "Weeks or months", "Fast (prompting)", "Fast (minutes)"),
        ("Live Speed", "Fast (<10 ms)", "Very slow (1-5 sec)", "Ultra-fast (<5 ms)"),
        ("Reliability", "100% consistent", "Unpredictable (hallucinates)", "100% consistent"),
        ("Customer Data Privacy", "Safe", "Risky (sent to cloud AI)", "100% safe (zero data to AI)"),
        ("Cost Per Transaction", "Free", "Expensive token fees", "Free (runs on local CPU)"),
        ("Handling API Changes", "Breaks silently", "Unpredictable", "Detects & alerts instantly")
    ]
    for r_idx, r_data in enumerate(comp_rows):
        bg = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = comp_table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
            if c_idx == 3:
                r.font.bold = True

    doc.add_paragraph()

    # 6. Key Patent Takeaways (What We Are Claiming)
    add_heading_1("6. Key Patent Takeaways (What We Are Claiming)")
    p_intro = doc.add_paragraph()
    r_intro = p_intro.add_run("These are the 5 core claims written in plain, straightforward terms so anyone can explain them:")
    r_intro.font.italic = True

    claims = [
        ("1. The 5-Tier Residual Matching Cascade",
         "A method that maps data fields between two systems by checking exact matches first, then normalized names, then a domain synonym dictionary, then past approved rules — and delegates only the leftover unmapped fields to an AI model with strict schema constraints, rather than wasting cost and risking errors on the whole dataset."),
        ("2. Decoupled AI-Advisory Rule Synthesis with Zero-AI Runtime",
         "An architecture that isolates generative AI strictly to an offline design-time assistant for discovering mapping rules, while routing all live production transactions through a pure, in-memory execution engine (<5ms) that never sends customer data to an AI model, eliminating latency, hallucination, and privacy risks."),
        ("3. Closed-Loop Cryptographic Schema Drift & Impact Mapping",
         "A monitoring system that fingerprints schemas using cryptographic hashes (SHA-256), detects additions, removals, or type changes via structural tree comparison, and automatically traces the change directly to affected transformation rules (BROKEN_PATH), alerting operators before broken data hits production."),
        ("4. Confidence-Gated Human-in-the-Loop Safeguard",
         "A verification gate that automatically scores every suggested mapping rule, allowing high-confidence rules to proceed while quarantining ambiguous or low-confidence rules into a review queue, preventing unverified or incorrect translations from entering live systems."),
        ("5. Automated Bi-Directional Schema Inversion",
         "A complete bi-directional integration system that takes forward transformation rules (such as legacy write into modern cloud) and automatically synthesizes the complementary reverse rules (such as modern cloud read back to legacy format), including automatic nesting, flattening, and format reversal without writing rules twice.")
    ]

    for title, desc in claims:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Inches(0.15)
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_t = p_c.add_run(title + ":")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
        p_d = doc.add_paragraph()
        p_d.paragraph_format.left_indent = Inches(0.3)
        p_d.paragraph_format.space_after = Pt(6)
        p_d.add_run(desc)

    # 7. Next Steps to File
    add_heading_1("7. Simple Next Steps to File")
    add_bullet("Add your name and company name as the inventor.", "1. Add Your Details")
    add_bullet("File this document as a Provisional Patent Application. It is simple, fast, inexpensive, and locks in your priority date immediately.", "2. File a Provisional Patent")
    add_bullet("You have 12 full months from your filing date to finalize it with a patent attorney while you build and demo your product.", "3. 12 Months to Convert")

    output_path = r"c:\Users\baljinders\Documents\GitHub\SchemaBridgeAI\Patent_Invention_Disclosure_SchemaBridgeAI.docx"
    doc.save(output_path)
    print(f"Document saved successfully to {output_path}")

if __name__ == "__main__":
    create_simple_patent_doc()
